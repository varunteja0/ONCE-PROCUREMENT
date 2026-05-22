"""L3.7 — Streaming XLSX parser (openpyxl read-only mode)."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

__all__ = ["iter_xlsx_rows"]


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


def _normalize_header(name: Any) -> str:
    if name is None:
        return ""
    return str(name).replace("\ufeff", "").strip().lower()


def iter_xlsx_rows(path: Path) -> Iterator[dict[str, str]]:
    """Yield one ``{header: value}`` dict per data row of sheet 0.

    Workbook is opened ``read_only=True, data_only=True`` so memory stays
    bounded regardless of file size and formulas surface as their cached
    values (the raw formula text would be useless to an importer).
    """

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        if not wb.sheetnames:
            return
        ws = wb[wb.sheetnames[0]]
        row_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(row_iter)
        except StopIteration:
            return
        headers = [_normalize_header(c) for c in header_row]
        width = len(headers)
        for raw in row_iter:
            if raw is None:
                continue
            cells = list(raw)
            if len(cells) < width:
                cells.extend([None] * (width - len(cells)))
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
                continue
            out: dict[str, str] = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                out[header] = _stringify(cells[idx] if idx < len(cells) else None)
            yield out
    finally:
        wb.close()
